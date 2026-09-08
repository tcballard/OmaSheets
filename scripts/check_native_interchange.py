#!/usr/bin/env python3
"""Verify native presentation and XLSX conversion against an independent reader."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import time


def main() -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    parser = argparse.ArgumentParser()
    parser.add_argument("--service", type=Path, required=True)
    parser.add_argument("--directory", type=Path, required=True)
    args = parser.parse_args()
    directory = args.directory.resolve()
    directory.mkdir(parents=True, exist_ok=False)
    runtime = directory / "runtime"
    runtime.mkdir(mode=0o700)
    native, source, exported = (directory / name for name in ("controls.omasheets", "source.xlsx", "export.xlsx"))
    service = str(args.service.resolve())
    actor = {"kind": "human", "id": "interchange-check"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    for row in [["Quarterly plan"], ["Region", "Budget", "Actual", "Variance"],
                ["North", 120, 95, "=B3-C3"], ["South", 80, 85, "=B4-C4"],
                ["Total", "=SUM(B3:B4)", "=SUM(C3:C4)", "=B5-C5"]]:
        ws.append(row)
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["A"].width = 25
    ws.freeze_panes = "C3"
    ws.sheet_view.showGridLines = False
    for cell in ws[2]:
        cell.font = Font(bold=True, size=12, color="FFE1E2E7")
        cell.fill = PatternFill("solid", fgColor="FF2D3D59")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin"))
    ws["B3"].number_format = "£#,##0.00"
    ws["F9"].font = Font(italic=True, underline="single", size=17, color="FF112233")
    ws["F9"].fill = PatternFill("solid", fgColor="FFEEDDCC")
    ws["B4"].comment = Comment("Check this estimate", "Workbook author")
    wb.save(source)
    original = hashlib.sha256(source.read_bytes()).hexdigest()
    with (directory / "service.log").open("w") as log:
        process = subprocess.Popen([service, "serve", "--runtime-dir", str(runtime)], stdout=log, stderr=log)
        try:
            for _ in range(150):
                if (runtime / "omasheets/native.sock").exists():
                    break
                if process.poll() is not None:
                    raise RuntimeError("Service exited during startup")
                time.sleep(0.02)

            def call(kind: str, *, success=True, **arguments):
                result = subprocess.run([service, "call", "--runtime-dir", str(runtime),
                                         json.dumps({"kind": kind, **arguments})],
                                        text=True, capture_output=True, timeout=35)
                answer = json.loads(result.stdout)
                assert answer["ok"] == success, answer
                return answer["response"] if success else answer

            def doc(kind: str, **arguments):
                return call(kind, path=str(native), **arguments)

            imported = call("import_xlsx", source=str(source), output=str(native), actor=actor, name=None)
            assert any("comments" in item for item in imported["limitations"])
            assert hashlib.sha256(source.read_bytes()).hexdigest() == original
            sheet = doc("document")["sheets"][0]["id"]
            assert doc("cell", sheet=sheet, a1="D3")["value"]["value"] == 25
            view = doc("sheet_view", sheet=sheet)
            assert (view["frozen_rows"], view["frozen_columns"]) == (2, 2)
            assert view["show_grid_lines"] is False and view["merges"][0]["columns"] == 4
            page = doc("grid_page", sheet=sheet, row_start=8, column_start=5, rows=1, columns=1)
            assert page["cells"][0]["style"]["italic"] is True
            revision = doc("revision")["revision"]
            changed = doc("edit_sheet", sheet=sheet, expected_revision=revision,
                          action={"action": "chart", "title": "Budget and actual", "kind": "bar",
                                  "range": {"row": 1, "column": 0, "rows": 3, "columns": 3}})
            chart = doc("sheet_view", sheet=sheet)["charts"][0]
            assert chart["series"][0]["values"] == [120, 80]
            assert changed["undo"]
            digest = doc("document")["digest"]
            doc("close")
            assert doc("document")["digest"] == digest
            manifest = doc("export_xlsx", output=str(exported))
            assert manifest["formula_cells_flattened"] == 0
            independent = load_workbook(exported)
            actual = independent["Plan"]
            assert actual["D3"].value == "=B3-C3"
            assert actual["B3"].number_format == "£#,##0.00"
            assert actual["A2"].font.bold and actual["A2"].font.sz == 12
            assert actual["A2"].fill.fgColor.rgb.upper() == "FF2D3D59"
            assert actual["A2"].alignment.horizontal == "center" and actual["A2"].alignment.wrap_text
            assert actual["A2"].border.bottom.style == "thin"
            assert actual["F9"].font.italic and actual["F9"].font.underline == "single"
            assert actual["F9"].fill.fgColor.rgb.upper() == "FFEEDDCC"
            assert actual.row_dimensions[1].height == 24
            assert abs(actual.column_dimensions["A"].width - 25) < 0.01
            assert actual.freeze_panes == "C3" and str(actual.merged_cells) == "A1:D1"
            assert actual.sheet_view.showGridLines is False
            assert load_workbook(exported, data_only=True)["Plan"]["D3"].value == 25
            exported_hash = hashlib.sha256(exported.read_bytes()).hexdigest()
            call("export_xlsx", success=False, path=str(native), output=str(exported))
            call("import_xlsx", success=False, source=str(source), output=str(native), actor=actor, name=None)
            assert hashlib.sha256(exported.read_bytes()).hexdigest() == exported_hash
            assert hashlib.sha256(source.read_bytes()).hexdigest() == original
            report = {"status": "ok", "source_sha256": original, "digest": digest,
                      "checks": ["independent styled import/export", "blank-cell style", "formulas and cached values",
                                 "dimensions", "merges", "frozen panes", "chart values", "close/reopen", "no overwrite"],
                      "limitations": manifest["limitations"]}
            (directory / "report.json").write_text(json.dumps(report, indent=2) + "\n")
            print(json.dumps(report))
        finally:
            process.terminate()
            process.wait(timeout=10)


if __name__ == "__main__":
    main()
